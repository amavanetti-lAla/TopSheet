"""Logica senza interfaccia: lettura del topsheet xlsx e generazione del PDF."""
import io
import pandas as pd
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

COLS = ["id", "brand", "stagione", "cliente", "trade_name", "city", "qty", "pct", "sel", "oc", "prod",
        "payment_terms", "producer_code", "rc_del", "mje_del", "comments",
        "pay_status", "pay_paid", "pay_due", "ship_status", "ship_date", "updated_at", "updated_by"]
PAY = ["da_pagare", "acconto_pagato", "saldato"]
SHIP = ["da_spedire", "spedito", "consegnato"]


def num(x):
    if x is None or (isinstance(x, float) and pd.isna(x)) or x == "":
        return ""
    return str(int(x)) if float(x).is_integer() else str(x)


def eur(x):
    """Importo in stile italiano: €17.249,00 (vuoto se non è un numero)."""
    try:
        v = float(x)
    except (TypeError, ValueError):
        return ""
    if pd.isna(v):
        return ""
    return "€" + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def parse_xlsx(f):
    """Legge un topsheet: titolo in A1 ('BRAND STAGIONE'), righe tra '... ACCOUNTS' e 'GRAND TOTALS'."""
    ws = load_workbook(f, data_only=True).active
    brand, _, season = str(ws["A1"].value or "").strip().rpartition(" ")
    first = next((r for r in range(1, ws.max_row + 1)
                  if "ACCOUNTS" in str(ws.cell(r, 1).value or "").upper()), None)
    if first is None:
        raise ValueError("Non trovo la riga 'EUR ACCOUNTS': il formato del foglio non è riconosciuto.")
    rows = []
    for r in range(first + 1, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 15)]
        if str(v[0] or "").upper().startswith("GRAND TOTALS"):
            break
        if not v[1]:
            continue
        s = lambda x: "" if x is None else str(x).strip()
        y = lambda x: s(x).upper() == "Y"
        terms = s(v[9])
        rows.append(dict(cliente=s(v[0]), trade_name=s(v[1]), city=s(v[2]), qty=v[3], pct=v[4],
                         sel=y(v[5]), oc=y(v[7]), prod=y(v[8]), payment_terms=terms,
                         producer_code=s(v[10]), rc_del=s(v[11]), mje_del=s(v[12]), comments=s(v[13]),
                         pay_status="acconto_pagato" if "DP PAID" in terms.upper() else "da_pagare",
                         ship_status="da_spedire"))
    return brand.strip(), season.strip(), pd.DataFrame(rows)


GREEN, YEL, PEACH = (colors.HexColor(h) for h in ("#E2EFDA", "#FFFF99", "#FCE4D6"))
W = [18.9, 18, 7.4, 3.5, 11.4, 3.4, 5.75, 3.1, 5.25, 14.4, 4.5, 14.6, 17, 23.9]  # larghezze colonne del foglio


def make_pdf(df, brand, season):
    """PDF A4 orizzontale con la stessa struttura del topsheet Excel."""
    page, m = landscape(A4), 18
    k = (page[0] - 2 * m) / sum(W)
    b = ParagraphStyle("b", fontName="Helvetica-Bold", fontSize=6.5, leading=7.5, alignment=1)
    l = ParagraphStyle("l", parent=b, fontName="Helvetica", alignment=0)
    big = ParagraphStyle("t", parent=b, fontSize=11, leading=13)
    P = lambda t, s=b: Paragraph(escape(str(t)), s)
    E = ""
    rows = [
        [P(f"{brand} {season}", big)] + [E] * 9 + [P("COMMENTS", big)] + [E] * 3,
        [E] * 3 + [P(season), E, P("SEL"), E, P("OC"), P("PROD"), P("PAYMENT TERMS")] + [E] * 4,
        [E] * 14,
        [P("SHOP NAME"), P("TRADE NAME"), P("CITY")] + [E] * 11,
        [P("EUR ACCOUNTS")] + [E] * 9 + [P(season), P("RC DEL"), P("MJE DEL"), P("COMMENTS")],
    ]
    yes = lambda x: P("Y") if x else E
    for r in df.itertuples():
        rows.append([P(r.cliente), P(r.trade_name), P(r.city), P(num(r.qty)), P(eur(r.pct)), yes(r.sel), E,
                     yes(r.oc), yes(r.prod), P(r.payment_terms), P(r.producer_code), P(r.rc_del), P(r.mje_del),
                     P(r.comments, l)])
    q = pd.to_numeric(df["qty"], errors="coerce").sum()
    p = pd.to_numeric(df["pct"], errors="coerce").sum()
    rows.append([P("GRAND TOTALS EUR"), E, E, P(num(q)), P(eur(p)), E, P("0")] + [E] * 7)
    t, last = len(rows) - 1, len(rows) - 1
    buf = io.BytesIO()
    tbl = Table(rows, colWidths=[w * k for w in W], repeatRows=5)
    tbl.setStyle(TableStyle([
        ("SPAN", (0, 0), (2, 2)), ("SPAN", (3, 0), (4, 0)), ("SPAN", (3, 1), (4, 1)), ("SPAN", (5, 1), (6, 1)),
        ("SPAN", (10, 0), (13, 2)), ("SPAN", (0, 4), (9, 4)), ("SPAN", (0, last), (2, last)),
        ("BACKGROUND", (0, 0), (13, 2), GREEN), ("BACKGROUND", (0, 3), (2, 3), YEL),
        ("BACKGROUND", (0, 4), (13, 4), PEACH), ("BACKGROUND", (0, last), (13, last), PEACH),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
    ]))
    SimpleDocTemplate(buf, pagesize=page, leftMargin=m, rightMargin=m, topMargin=m, bottomMargin=m,
                      title=f"{brand} {season} topsheet").build([tbl])
    return buf.getvalue()
